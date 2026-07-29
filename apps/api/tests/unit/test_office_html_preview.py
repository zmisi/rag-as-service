"""Unit tests for Office → HTML preview conversion."""

from __future__ import annotations

import io

from docx import Document
from openpyxl import Workbook
from pptx import Presentation

from rag_api.services.office_html_preview import office_to_preview_html


def _docx_bytes(text: str) -> bytes:
    doc = Document()
    doc.add_heading("第二章 索引优化策略", level=1)
    doc.add_heading("2.1 索引类型与选择", level=2)
    doc.add_paragraph(text)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "索引类型"
    table.cell(0, 1).text = "使用范围"
    table.cell(1, 0).text = "B-Tree"
    table.cell(1, 1).text = "等值与范围查询"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _xlsx_bytes(phrase: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "产品信息"
    ws.append(["产品线", "产品型号", "规格参数", "价格(元)"])
    ws.append(["智能手机", "X100 Pro", "6.8寸 OLED", 6999])
    ws.append([None, "X100", "6.7寸 AMOLED", 4999])
    ws.append([None, "X90 Pro", "6.78寸 AMOLED", 5999])
    ws.merge_cells("A2:A4")
    ws["B2"] = "X100 Pro"
    ws["C2"] = f"6.8寸 OLED, {phrase}"
    ws2 = wb.create_sheet("技术参数表")
    ws2.append(["产品型号", "处理器"])
    ws2.append(["X100 Pro", "骁龙8 Gen 3"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pptx_bytes(phrase: str) -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "索引策略"
    slide.placeholders[1].text = phrase
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_docx_preview_is_semantic_html_not_markdown():
    html = office_to_preview_html("guide.docx", _docx_bytes("B-Tree索引适用场景"))
    assert "<!DOCTYPE html>" in html
    assert "<h1>" in html or "<h2>" in html
    assert "B-Tree索引适用场景" in html
    assert "<table>" in html
    assert "| 索引类型 |" not in html
    assert "<pre>" not in html


def test_xlsx_preview_renders_excel_like_grid():
    html = office_to_preview_html("sheet.xlsx", _xlsx_bytes("UNIQUE_XLSX_PREVIEW"))
    assert 'class="xlsx-body"' in html
    assert 'class="xlsx-grid"' in html
    assert 'class="col-head">A</th>' in html
    assert 'class="row-head">1</th>' in html
    assert "产品信息" in html
    assert "技术参数表" in html
    assert "UNIQUE_XLSX_PREVIEW" in html
    assert 'rowspan="3"' in html
    assert 'data-sheet="sheet-0"' in html
    assert 'data-sheet="sheet-1"' in html
    assert "| 项目 |" not in html


def test_pptx_preview_renders_slide_cards():
    html = office_to_preview_html("deck.pptx", _pptx_bytes("UNIQUE_PPTX_PREVIEW"))
    assert "幻灯片 1" in html
    assert "索引策略" in html
    assert "UNIQUE_PPTX_PREVIEW" in html
    assert 'class="slide-card"' in html

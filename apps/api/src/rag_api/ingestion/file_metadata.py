"""Industry-style file_metadata for documents (schema_version=1).

upload: provenance at HTTP upload time
document: properties extracted at ingest (PDF/Office); merge, never fail the job
"""

from __future__ import annotations

import io
import logging
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

logger = logging.getLogger(__name__)

SCHEMA_VERSION = 1


def _utc_now_iso() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace(
        "+00:00", "Z"
    )


def _omit_empty(d: dict[str, Any]) -> dict[str, Any]:
    out: dict[str, Any] = {}
    for k, v in d.items():
        if v is None:
            continue
        if isinstance(v, str) and not v.strip():
            continue
        out[k] = v
    return out


def build_upload_metadata(
    *,
    filename: str,
    content_type: str,
    size_bytes: int,
    uploaded_at: str | None = None,
) -> dict[str, Any]:
    return {
        "schema_version": SCHEMA_VERSION,
        "upload": _omit_empty(
            {
                "uploaded_at": uploaded_at or _utc_now_iso(),
                "original_filename": filename,
                "declared_content_type": content_type or "application/octet-stream",
                "size_bytes": int(size_bytes),
            }
        ),
    }


def merge_file_metadata(
    existing: dict[str, Any] | None,
    *,
    upload: dict[str, Any] | None = None,
    document: dict[str, Any] | None = None,
) -> dict[str, Any]:
    base: dict[str, Any] = dict(existing or {})
    base["schema_version"] = SCHEMA_VERSION
    if upload is not None:
        prev = dict(base.get("upload") or {})
        prev.update(_omit_empty(upload))
        base["upload"] = prev
    if document is not None:
        prev = dict(base.get("document") or {})
        prev.update(_omit_empty(document))
        if prev:
            base["document"] = prev
    return base


def parse_iso_to_naive_utc(value: str | None) -> datetime | None:
    if not value or not isinstance(value, str):
        return None
    s = value.strip()
    if not s:
        return None
    try:
        if s.endswith("Z"):
            s = s[:-1] + "+00:00"
        dt = datetime.fromisoformat(s)
        if dt.tzinfo is not None:
            dt = dt.astimezone(timezone.utc).replace(tzinfo=None)
        return dt
    except ValueError:
        return None


def _pdf_properties(data: bytes) -> dict[str, Any]:
    import fitz  # PyMuPDF

    doc = fitz.open(stream=data, filetype="pdf")
    try:
        meta = doc.metadata or {}
        props = _omit_empty(
            {
                "title": meta.get("title"),
                "author": meta.get("author"),
                "creator": meta.get("creator"),
                "subject": meta.get("subject"),
                "created_at": meta.get("creationDate") or meta.get("created"),
                "modified_at": meta.get("modDate") or meta.get("moddate"),
                "page_count": int(doc.page_count) if doc.page_count else None,
            }
        )
        return props
    finally:
        doc.close()


def _docx_properties(data: bytes) -> dict[str, Any]:
    from docx import Document

    doc = Document(io.BytesIO(data))
    cp = doc.core_properties
    props = _omit_empty(
        {
            "title": cp.title,
            "author": cp.author,
            "creator": cp.author,
            "subject": cp.subject,
            "created_at": cp.created.isoformat() + "Z" if cp.created else None,
            "modified_at": cp.modified.isoformat() + "Z" if cp.modified else None,
        }
    )
    return props


def _pptx_properties(data: bytes) -> dict[str, Any]:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    cp = prs.core_properties
    props = _omit_empty(
        {
            "title": cp.title,
            "author": cp.author,
            "creator": cp.author,
            "subject": cp.subject,
            "created_at": cp.created.isoformat() + "Z" if cp.created else None,
            "modified_at": cp.modified.isoformat() + "Z" if cp.modified else None,
            "slide_count": len(prs.slides),
        }
    )
    return props


def _xlsx_properties(data: bytes) -> dict[str, Any]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        props: dict[str, Any] = {"sheet_count": len(wb.sheetnames)}
        # openpyxl props live on wb.properties
        p = wb.properties
        if p is not None:
            props.update(
                _omit_empty(
                    {
                        "title": p.title,
                        "author": p.creator,
                        "creator": p.creator,
                        "subject": p.subject,
                        "created_at": p.created.isoformat() + "Z" if p.created else None,
                        "modified_at": p.modified.isoformat() + "Z"
                        if p.modified
                        else None,
                    }
                )
            )
        return _omit_empty(props)
    finally:
        wb.close()


def _text_charset_hint(data: bytes) -> dict[str, Any]:
    for enc in ("utf-8", "utf-8-sig", "gb18030", "latin-1"):
        try:
            data.decode(enc)
            return {"charset": enc}
        except UnicodeDecodeError:
            continue
    return {}


def extract_document_properties(filename: str, data: bytes) -> dict[str, Any]:
    """Best-effort document properties; never raises to caller for known formats."""
    suffix = Path(filename).suffix.lower()
    try:
        if suffix == ".pdf":
            return _pdf_properties(data)
        if suffix == ".docx":
            return _docx_properties(data)
        if suffix == ".pptx":
            return _pptx_properties(data)
        if suffix == ".xlsx":
            return _xlsx_properties(data)
        if suffix in {".txt", ".md"}:
            return _text_charset_hint(data)
    except Exception:  # noqa: BLE001 — metadata must not break ingest
        logger.exception("extract_document_properties failed filename=%s", filename)
    return {}


def apply_modified_at_from_metadata(
    file_metadata: dict[str, Any] | None,
) -> datetime | None:
    """Prefer document.modified_at, else upload.uploaded_at."""
    meta = file_metadata or {}
    document = meta.get("document") or {}
    upload = meta.get("upload") or {}
    return parse_iso_to_naive_utc(document.get("modified_at")) or parse_iso_to_naive_utc(
        upload.get("uploaded_at")
    )

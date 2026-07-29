"""Office OOXML → read-only HTML preview (Admin).

Industry pattern: convert to semantic HTML for iframe preview.
docx uses mammoth; xlsx/pptx use lightweight HTML layouts.
"""

from __future__ import annotations

import datetime as dt
import html
import io
from pathlib import Path

from rag_api.ingestion.office import OfficeParseError

_OFFICE_SUFFIXES = frozenset({".docx", ".pptx", ".xlsx"})

_PREVIEW_CSS = """
:root { color-scheme: light; }
* { box-sizing: border-box; }
body {
  margin: 0;
  padding: 1.25rem 1.5rem 2rem;
  font: 15px/1.6 system-ui, -apple-system, "Segoe UI", sans-serif;
  color: #1a1d21;
  background: #f4f6f8;
}
.doc-shell {
  max-width: 52rem;
  margin: 0 auto;
  background: #fff;
  border: 1px solid #e2e6ea;
  border-radius: 10px;
  padding: 1.5rem 1.75rem 2rem;
  box-shadow: 0 1px 2px rgba(16, 24, 40, 0.04);
}
.doc-shell h1, .doc-shell h2, .doc-shell h3,
.doc-shell h4, .doc-shell h5, .doc-shell h6 {
  line-height: 1.3;
  margin: 1.25em 0 0.5em;
  color: #101828;
}
.doc-shell h1 { font-size: 1.55rem; margin-top: 0; }
.doc-shell h2 { font-size: 1.28rem; }
.doc-shell h3 { font-size: 1.1rem; }
.doc-shell p { margin: 0.65em 0; }
.doc-shell ul, .doc-shell ol { margin: 0.65em 0; padding-left: 1.4em; }
.doc-shell table {
  border-collapse: collapse;
  width: 100%;
  margin: 1em 0;
  font-size: 0.92rem;
}
.doc-shell th, .doc-shell td {
  border: 1px solid #d0d5dd;
  padding: 0.45rem 0.65rem;
  text-align: left;
  vertical-align: top;
}
.doc-shell th { background: #f8fafc; font-weight: 650; }
.doc-shell img { max-width: 100%; height: auto; }
body.xlsx-body {
  padding: 0;
  background: #e8eaed;
  height: 100vh;
  overflow: hidden;
}
.xlsx-viewer {
  display: flex;
  flex-direction: column;
  height: 100vh;
  background: #fff;
}
.xlsx-toolbar {
  flex: 0 0 auto;
  padding: 0.55rem 0.85rem;
  border-bottom: 1px solid #c7c7c7;
  background: #f3f3f3;
  color: #333;
  font-size: 0.82rem;
}
.xlsx-scroll {
  flex: 1 1 auto;
  overflow: auto;
  background: #fff;
}
.xlsx-sheet { display: none; }
.xlsx-sheet.active { display: block; }
.xlsx-grid {
  border-collapse: collapse;
  table-layout: fixed;
  background: #fff;
  font: 13px/1.35 "Segoe UI", system-ui, sans-serif;
  color: #000;
}
.xlsx-grid .corner,
.xlsx-grid .col-head,
.xlsx-grid .row-head {
  background: #f8f9fa;
  color: #5f6368;
  border: 1px solid #c0c0c0;
  text-align: center;
  font-weight: 600;
  user-select: none;
}
.xlsx-grid .corner {
  position: sticky;
  left: 0;
  top: 0;
  z-index: 3;
  width: 3rem;
  min-width: 3rem;
}
.xlsx-grid .col-head {
  position: sticky;
  top: 0;
  z-index: 2;
  min-width: 6.5rem;
  padding: 0.2rem 0.35rem;
}
.xlsx-grid .row-head {
  position: sticky;
  left: 0;
  z-index: 1;
  width: 3rem;
  min-width: 3rem;
  padding: 0.15rem 0.25rem;
}
.xlsx-grid td.cell {
  border: 1px solid #d0d7de;
  padding: 0.18rem 0.45rem;
  vertical-align: middle;
  white-space: pre-wrap;
  word-break: break-word;
  min-width: 6.5rem;
  max-width: 18rem;
  background: #fff;
}
.xlsx-grid td.cell.num { text-align: right; font-variant-numeric: tabular-nums; }
.xlsx-grid td.cell.center { text-align: center; }
.xlsx-grid td.cell.bold { font-weight: 700; }
.xlsx-tabs {
  flex: 0 0 auto;
  display: flex;
  gap: 0;
  align-items: flex-end;
  padding: 0 0.35rem;
  border-top: 1px solid #c7c7c7;
  background: #f3f3f3;
  overflow-x: auto;
  min-height: 2.1rem;
}
.xlsx-tab {
  appearance: none;
  border: 1px solid transparent;
  border-bottom: none;
  border-radius: 0;
  background: transparent;
  color: #333;
  font: inherit;
  font-size: 0.8rem;
  padding: 0.35rem 0.9rem;
  cursor: pointer;
  white-space: nowrap;
}
.xlsx-tab:hover { background: #e8eaed; }
.xlsx-tab.active {
  background: #fff;
  border-color: #c7c7c7;
  border-top: 2px solid #217346;
  font-weight: 650;
}
.slide-stack { display: grid; gap: 1rem; }
.slide-card {
  background: #fff;
  border: 1px solid #e2e6ea;
  border-radius: 10px;
  padding: 1.1rem 1.25rem 1.25rem;
  box-shadow: 0 1px 2px rgba(16, 24, 40, 0.04);
}
.slide-label {
  margin: 0 0 0.55rem;
  font-size: 0.72rem;
  font-weight: 700;
  letter-spacing: 0.04em;
  text-transform: uppercase;
  color: #667085;
}
.slide-card h2 {
  margin: 0 0 0.55rem;
  font-size: 1.15rem;
}
.slide-card p { margin: 0.4em 0; white-space: pre-wrap; }
.slide-notes {
  margin-top: 0.85rem;
  padding-top: 0.75rem;
  border-top: 1px dashed #e2e6ea;
  color: #475467;
  font-size: 0.9rem;
}
.empty {
  color: #667085;
  font-style: italic;
}
""".strip()


def office_to_preview_html(filename: str, data: bytes) -> str:
    """Convert OOXML bytes to a full read-only HTML preview document."""
    suffix = Path(filename).suffix.lower()
    if suffix not in _OFFICE_SUFFIXES:
        raise OfficeParseError(f"Not an OOXML office type: {suffix}")
    body_class = ""
    if suffix == ".docx":
        body = _docx_to_html_fragment(data)
        fallback = '<p class="empty">(空文档)</p>'
        inner = f'<article class="doc-shell">{body or fallback}</article>'
    elif suffix == ".xlsx":
        body_class = "xlsx-body"
        inner = _xlsx_to_html_fragment(data)
    else:
        inner = _pptx_to_html_fragment(data)
    return _wrap_preview_document(filename, inner, body_class=body_class)


def _wrap_preview_document(
    filename: str,
    body_html: str,
    *,
    body_class: str = "",
) -> str:
    """Wrap an HTML fragment in a styled standalone preview page."""
    title = html.escape(filename or "preview")
    cls = f' class="{html.escape(body_class)}"' if body_class else ""
    return (
        "<!DOCTYPE html><html lang=\"zh-CN\"><head>"
        '<meta charset="utf-8"/>'
        f"<title>{title}</title>"
        f"<style>{_PREVIEW_CSS}</style>"
        f"</head><body{cls}>"
        f"{body_html}"
        "</body></html>"
    )


def _docx_to_html_fragment(data: bytes) -> str:
    """Convert .docx to semantic HTML via mammoth (industry-standard path)."""
    try:
        import mammoth
    except ImportError as exc:
        raise OfficeParseError(
            "mammoth is not installed; pip install mammoth"
        ) from exc

    try:
        result = mammoth.convert_to_html(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise OfficeParseError(f"Failed to convert .docx to HTML: {exc}") from exc
    return (result.value or "").strip()


def _xlsx_to_html_fragment(data: bytes) -> str:
    """Render .xlsx as an Excel-like grid with sheet tabs and merges."""
    try:
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise OfficeParseError(
            "openpyxl is not installed; pip install openpyxl"
        ) from exc

    try:
        # Non-read-only so merged cell ranges are available.
        wb = load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise OfficeParseError(f"Failed to open .xlsx: {exc}") from exc

    try:
        sheet_panels: list[str] = []
        tab_buttons: list[str] = []
        for idx, sheet in enumerate(wb.worksheets):
            sheet_id = f"sheet-{idx}"
            active = " active" if idx == 0 else ""
            title = sheet.title or f"Sheet{idx + 1}"
            tab_buttons.append(
                "<button type=\"button\" class=\"xlsx-tab"
                f'{active}" data-sheet="{sheet_id}">'
                f"{html.escape(title)}</button>"
            )
            grid = _xlsx_sheet_grid_html(sheet, get_column_letter)
            sheet_panels.append(
                f'<div class="xlsx-sheet{active}" id="{sheet_id}">{grid}</div>'
            )
    except Exception as exc:  # noqa: BLE001
        raise OfficeParseError(f"Failed to parse .xlsx content: {exc}") from exc
    finally:
        wb.close()

    script = """
<script>
(function () {
  var tabs = document.querySelectorAll(".xlsx-tab");
  var sheets = document.querySelectorAll(".xlsx-sheet");
  tabs.forEach(function (tab) {
    tab.addEventListener("click", function () {
      var id = tab.getAttribute("data-sheet");
      tabs.forEach(function (t) { t.classList.toggle("active", t === tab); });
      sheets.forEach(function (s) {
        s.classList.toggle("active", s.id === id);
      });
    });
  });
})();
</script>
""".strip()

    return (
        '<div class="xlsx-viewer">'
        '<div class="xlsx-toolbar">只读预览 · 按原表格布局展示</div>'
        f'<div class="xlsx-scroll">{"".join(sheet_panels)}</div>'
        f'<div class="xlsx-tabs">{"".join(tab_buttons)}</div>'
        f"{script}"
        "</div>"
    )


_MAX_PREVIEW_ROWS = 200
_MAX_PREVIEW_COLS = 40


def _xlsx_sheet_grid_html(sheet: object, get_column_letter: object) -> str:
    """Build one worksheet as an Excel-like HTML grid."""
    max_row = int(getattr(sheet, "max_row", 0) or 0)
    max_col = int(getattr(sheet, "max_column", 0) or 0)
    if max_row <= 0 or max_col <= 0:
        return '<p class="empty" style="padding:1rem">(空工作表)</p>'

    max_row = min(max_row, _MAX_PREVIEW_ROWS)
    max_col = min(max_col, _MAX_PREVIEW_COLS)

    merge_origins, merge_covered = _xlsx_merge_info(sheet)

    # Column header row: corner + A B C ...
    head_cells = ['<th class="corner"></th>']
    for col in range(1, max_col + 1):
        letter = get_column_letter(col)
        head_cells.append(f'<th class="col-head">{html.escape(str(letter))}</th>')
    rows_html = [f"<tr>{''.join(head_cells)}</tr>"]

    for row in range(1, max_row + 1):
        cells = [f'<th class="row-head">{row}</th>']
        col = 1
        while col <= max_col:
            key = (row, col)
            if key in merge_covered:
                col += 1
                continue

            rowspan = 1
            colspan = 1
            if key in merge_origins:
                rowspan, colspan = merge_origins[key]

            cell = sheet.cell(row=row, column=col)
            text = _preview_cell(cell.value)
            classes = ["cell"]
            if isinstance(cell.value, (int, float)) and not isinstance(
                cell.value, bool
            ):
                classes.append("num")
            align = getattr(getattr(cell, "alignment", None), "horizontal", None)
            if align == "center":
                classes.append("center")
            elif align == "right":
                classes.append("num")
            font = getattr(cell, "font", None)
            if font is not None and getattr(font, "bold", False):
                classes.append("bold")

            attrs = [f'class="{" ".join(classes)}"']
            if rowspan > 1:
                attrs.append(f'rowspan="{rowspan}"')
            if colspan > 1:
                attrs.append(f'colspan="{colspan}"')
            content = html.escape(text) if text else ""
            cells.append(f"<td {' '.join(attrs)}>{content}</td>")
            col += colspan
        rows_html.append(f"<tr>{''.join(cells)}</tr>")

    return (
        '<table class="xlsx-grid" role="grid">'
        f"<tbody>{''.join(rows_html)}</tbody>"
        "</table>"
    )


def _xlsx_merge_info(
    sheet: object,
) -> tuple[dict[tuple[int, int], tuple[int, int]], set[tuple[int, int]]]:
    """Return merge origins → (rowspan, colspan) and covered non-origin cells."""
    origins: dict[tuple[int, int], tuple[int, int]] = {}
    covered: set[tuple[int, int]] = set()
    ranges = getattr(getattr(sheet, "merged_cells", None), "ranges", None) or []
    for cell_range in ranges:
        min_row = int(cell_range.min_row)
        max_row = int(cell_range.max_row)
        min_col = int(cell_range.min_col)
        max_col = int(cell_range.max_col)
        origins[(min_row, min_col)] = (
            max_row - min_row + 1,
            max_col - min_col + 1,
        )
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                if (r, c) != (min_row, min_col):
                    covered.add((r, c))
    return origins, covered


def _pptx_to_html_fragment(data: bytes) -> str:
    """Render .pptx as a stack of read-only slide cards."""
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise OfficeParseError(
            "python-pptx is not installed; pip install python-pptx"
        ) from exc

    try:
        prs = Presentation(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise OfficeParseError(f"Failed to open .pptx: {exc}") from exc

    cards: list[str] = ['<div class="slide-stack">']
    try:
        for idx, slide in enumerate(prs.slides, start=1):
            texts: list[str] = []
            for shape in slide.shapes:
                if not getattr(shape, "has_text_frame", False):
                    continue
                chunk = (shape.text_frame.text or "").strip()
                if chunk:
                    texts.append(chunk)
            notes = ""
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes = (slide.notes_slide.notes_text_frame.text or "").strip()

            title = texts[0] if texts else f"幻灯片 {idx}"
            body_parts = texts[1:] if len(texts) > 1 else []
            body_html = "".join(
                f"<p>{html.escape(part)}</p>" for part in body_parts
            )
            if not body_html:
                body_html = '<p class="empty">(无正文)</p>'
            notes_html = (
                f'<div class="slide-notes"><strong>备注</strong>'
                f"<p>{html.escape(notes)}</p></div>"
                if notes
                else ""
            )
            cards.append(
                '<article class="slide-card">'
                f'<p class="slide-label">幻灯片 {idx}</p>'
                f"<h2>{html.escape(title)}</h2>"
                f"{body_html}"
                f"{notes_html}"
                "</article>"
            )
    except Exception as exc:  # noqa: BLE001
        raise OfficeParseError(f"Failed to parse .pptx content: {exc}") from exc

    cards.append("</div>")
    return "".join(cards)


def _preview_cell(value: object) -> str:
    """Format a spreadsheet cell for HTML preview (no Markdown escaping)."""
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        if (
            value.hour == 0
            and value.minute == 0
            and value.second == 0
            and value.microsecond == 0
        ):
            return value.date().isoformat()
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()
